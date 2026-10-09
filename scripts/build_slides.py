#!/usr/bin/env python3
"""Build a five-page business deck (PPTX + print-ready HTML) from shared content.

Usage: uv run python scripts/build_slides.py [--screenshot artifacts/map-overview.png]
       [--verification path/to/verification.json]
No network calls, authentication or model invocation.
"""

from __future__ import annotations

import argparse
import base64
from dataclasses import dataclass
from html import escape
import json
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt
    from PIL import Image as PILImage
except ImportError as exc:
    raise SystemExit("python-pptx (and its Pillow dependency) is required. Run the project's approved uv sync first.") from exc

from london_office.market import Market

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "artifacts" / "presentation"
MARKET = Market()
SOURCES = {
    "2024": MARKET.source("savills-2024-q4")["url"] + "#rents",
    "2025": MARKET.source("savills-2025-q4")["url"] + "#rents",
}
NAVY = "102A3A"
INK = "183446"
MUTED = "617582"
MINT = "B8E8D8"
TEAL = "206F69"
ORANGE = "BE7147"
PAPER = "F7F8F5"
WHITE = "FFFFFF"
BORDER = "DCE4E2"


def rent_story(area: str) -> tuple[str, str, str]:
    comparison = MARKET.compare(area=area, metric="prime_rent", previous_period="2024-YE", current_period="2025-YE")
    if comparison["status"] != "available":
        raise ValueError("Presentation example requires verified comparable source observations")
    return (f"£{comparison['previous_value']:.2f}", f"£{comparison['current_value']:.2f}", f"{comparison['change_value']:.2f}%")


@dataclass(frozen=True)
class Page:
    eyebrow: str
    title: str
    subtitle: str
    notes: str


def statuses(path: Path | None) -> dict[str, str]:
    result = {key: "Not yet verified" for key in ("offline_tests", "map_e2e", "live_model")}
    if path is None:
        return result
    obj = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(obj, dict):
        raise ValueError("Verification JSON must be an object")
    for key in result:
        if key in obj:
            value = obj[key]
            if value not in ("Not yet verified", "Verified", "Failed"):
                raise ValueError(f"{key}: expected Not yet verified, Verified or Failed")
            result[key] = value
    return result


def content(verification: dict[str, str]) -> list[dict]:
    common = ("Sources: Savills, 2024 year-end report (published 2025-01-31), "
              f"{SOURCES['2024']}; 2025 year-end report (published 2026-01-30), "
              f"{SOURCES['2025']}. Original unit £/sq ft; billing period not specified. "
              "Average prime achieved rent, not prime headline rent; historical year-end report values, "
              "not live prices or quarterly transaction averages. No building-level inference.")
    return [
        {"page": Page("DECISION SUPPORT  /  01", "From report PDFs to a defensible market conversation",
                      "London office submarkets · an evidence-led, interactive proof of concept",
                      common + " Business framing: a user needs to compare like with like and inspect a source before acting."),
         "problem": "The friction", "problem_text": "Research is scattered across dated reports; a number without its definition or source can mislead.",
         "value": "The useful answer", "value_text": "Select a market, ask a question, and see the reported value, Python-calculated change and source together.",
         "boundary": "The honest boundary", "boundary_text": "A historical snapshot is not a live quote. Missing coverage stays visible rather than becoming zero.",
         "footer": "FIRST RELEASE  ·  CITY + WEST END HISTORICAL EVIDENCE  ·  FOUR ILLUSTRATIVE MAP SELECTORS"},
        {"page": Page("USER JOURNEY  /  02", "A question starts on the map — and ends at evidence",
                      "A short workflow for a business reviewer, not a claim of building-level precision",
                      common + " The four hand-drawn polygons are navigation selectors only, not Savills report boundaries. "
                      "The interface uses deck.gl and MapLibre. A screenshot, when supplied, is a local application capture, "
                      "not an official map or report graphic."),
         "steps": [("01", "Choose a submarket", "Select City or West End; change metric or year-end period. The other two selectors can expose missing coverage."),
                   ("02", "Ask in context", "Ask what changed. The Python/Codex workflow is intended to request only approved market tools."),
                   ("03", "Inspect the trail", "Read the reported figures, calculated comparison, warnings and source links beside the explanation.")],
         "no_image": "WORKFLOW VIEW  ·  No screenshot supplied — this is a process diagram, not a simulated UI.",
         "image_caption": "Local application capture · OpenFreeMap · © OpenMapTiles · data from OpenStreetMap; project-made illustrative selectors."},
        {"page": Page("EVIDENCE IN PRACTICE  /  03", "One comparison. Two reported periods. One visible caveat.",
                      "Savills average prime achieved rent · reported year-end snapshots · £/sq ft",
                      common + " The West End report states 6.1% growth for 2025; the displayed snapshots "
                      "157.50 and 166.61 calculate to 5.78%. Cause (including any base/sample revision) is unknown; "
                      "do not say the report is wrong. The values and computations come from the same Python Market Interface as the app and Agent."),
         "rows": [("City", *rent_story("city")), ("West End", *rent_story("west_end"))],
         "warning": "West End: snapshot calculation 5.78%; the 2025 report publishes 6.1%. The reason for the difference is unknown — show both, do not reconcile by guesswork.",
         "source_line": "Savills 2024 year-end (31 Jan 2025)  ↗     Savills 2025 year-end (30 Jan 2026)  ↗"},
        {"page": Page("TRUST DESIGN  /  04", "Credibility is a product feature, not a footnote",
                      "Rules that make a persuasive answer inspectable",
                      common + " Approved read-only Python business tools: market_catalog, market_query, compare_periods, "
                      "market_brief, source_details, verify_source. Project-local MCP restricts the tools; no shell, "
                      "personal MCP, open web search or external agents. Initial runtime uses the Codex Python SDK "
                      "with a subscription, subject to independent login/model availability. Architecture intent is not a live test result."),
         "rules": [("01", "Arithmetic is code", "Python calculates changes; the model does not invent arithmetic or new numeric facts."),
                   ("02", "Definitions travel", "Prime achieved rent is not headline rent. Report year-end is not a quarterly average."),
                   ("03", "Unknown is explicit", "Missing is not zero. Conflicting growth stays visible; no silent overwrite."),
                   ("04", "Tools have a fence", "Only project read-only business tools via local MCP; each result carries evidence links.")],
         "bar": "HISTORICAL REPORTS, NOT LIVE MARKET DATA  /  ILLUSTRATIVE GEOMETRY, NOT REPORT BOUNDARIES"},
        {"page": Page("REVIEW & NEXT STEPS  /  05", "A focused PoC, with a clear route to independent review",
                      "What the evaluator can test — and what still needs live access",
                      common + " Offline tests and map/E2E tests must be reported separately from authenticated model testing. "
                      "A 'Verified' label is set only by an explicit verification JSON supplied after actual execution; "
                      "this generator does not run those tests. The evaluator uses their own authorized subscription/account "
                      "and dedicated Codex home outside this repo. No automatic API-key billing fallback."),
         "scope": [("IN SCOPE", "Historical City / West End snapshots, Python comparisons, evidence-backed Q&A and map navigation."),
                   ("NOT CLAIMED", "Live quotes, building-level metrics, exact market GIS, live news, or a second AI runtime."),
                   ("NEXT", "Review missing markets and indicator definitions with stakeholders before adding licensed or new sources.")],
         "route": [("1", "Rebuild", "uv sync → uv run python scripts/build_slides.py; follow the project run instructions for the app."),
                   ("2", "Walk through", "Select City / 2025 year-end / average prime rent; inspect the comparison and source. Then select Canary Wharf to show missing coverage."),
                   ("3", "Use own account", "Authenticate the dedicated Codex SDK subscription profile outside the repo; ask a contextual question and inspect executed tool evidence.")],
         "verification": verification},
    ]


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def rect(slide, x, y, w, h, fill=WHITE, line=None, radius=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                   Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.fill.background() if line is None else None
    if line is not None:
        shape.line.color.rgb = rgb(line)
    if radius:
        try:
            shape.adjustments[0] = 0.08
        except (AttributeError, IndexError):
            pass
    return shape


def text(slide, value, x, y, w, h, *, size=15, color=INK, bold=False, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.01)
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = frame.paragraphs[0]
    if align:
        p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def base(slide, page: Page, n: int):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(PAPER)
    rect(slide, 0, 0, 13.333, 0.12, NAVY)
    text(slide, "LONDON  /  OFFICE MONITOR", 0.65, 0.30, 6.5, 0.28, size=10, color=TEAL, bold=True)
    text(slide, page.eyebrow, 10.1, 0.30, 2.55, 0.28, size=10, color=MUTED, align=PP_ALIGN.RIGHT)
    text(slide, page.title, 0.65, 0.86, 12.0, 0.85, size=29, color=NAVY, bold=True)
    text(slide, page.subtitle, 0.67, 1.72, 11.9, 0.50, size=14, color=MUTED)
    rect(slide, 0.65, 7.08, 12.0, 0.012, BORDER)
    text(slide, "HISTORICAL  ·  EVIDENCE-LED  ·  POC", 0.66, 7.14, 7.7, 0.21, size=9, color=MUTED)
    text(slide, f"{n:02d}  /  05", 11.5, 7.14, 1.13, 0.21, size=9, color=MUTED, align=PP_ALIGN.RIGHT)
    slide.notes_slide.notes_text_frame.text = page.notes


def add_image(slide, path: Path, x: float, y: float, w: float, h: float):
    with PILImage.open(path) as image:
        iw, ih = image.size
    if iw <= 0 or ih <= 0:
        raise ValueError("Screenshot has invalid dimensions")
    factor = min(w / iw, h / ih)
    sw, sh = iw * factor, ih * factor
    slide.shapes.add_picture(str(path), Inches(x + (w - sw) / 2), Inches(y + (h - sh) / 2),
                             width=Inches(sw), height=Inches(sh))


def build_ppt(pages: list[dict], screenshot: Path | None, destination: Path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for i, data in enumerate(pages, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        base(slide, data["page"], i)
        if i == 1:
            for j, (title, body) in enumerate(((data["problem"], data["problem_text"]),
                                                (data["value"], data["value_text"]),
                                                (data["boundary"], data["boundary_text"]))):
                x = 0.66 + j * 4.08
                rect(slide, x, 2.58, 3.82, 3.37, WHITE, BORDER, True)
                rect(slide, x + 0.25, 2.87, 0.45, 0.08, ORANGE if j == 0 else TEAL)
                text(slide, f"0{j + 1}", x + 0.25, 3.16, 0.9, 0.46, size=23, color=TEAL)
                text(slide, title, x + 0.25, 3.78, 3.32, 0.45, size=19, color=NAVY, bold=True)
                text(slide, body, x + 0.25, 4.38, 3.26, 1.2, size=15)
            text(slide, data["footer"], 0.68, 6.34, 11.8, 0.35, size=11, color=TEAL, bold=True)
        elif i == 2:
            if screenshot:
                rect(slide, 6.15, 2.42, 6.5, 3.94, WHITE, BORDER, True)
                add_image(slide, screenshot, 6.30, 2.57, 6.2, 3.60)
                text(slide, data["image_caption"], 6.18, 6.48, 6.46, 0.42, size=9, color=MUTED)
            else:
                rect(slide, 6.23, 2.65, 6.1, 3.45, NAVY, radius=True)
                text(slide, "SELECT  →  ASK  →  CHECK", 6.75, 3.28, 5.1, 0.62, size=24, color=MINT, bold=True)
                text(slide, "Market context                 Python tools                 Source trail",
                     6.75, 4.13, 5.0, 0.65, size=13, color=WHITE)
                text(slide, data["no_image"], 6.27, 6.37, 6.02, 0.49, size=10, color=MUTED)
            for j, (num, title, body) in enumerate(data["steps"]):
                y = 2.47 + j * 1.46
                text(slide, num, 0.72, y, 0.73, 0.51, size=22, color=TEAL, bold=True)
                text(slide, title, 1.49, y, 4.47, 0.39, size=17, color=NAVY, bold=True)
                text(slide, body, 1.49, y + 0.45, 4.4, 0.82, size=12)
        elif i == 3:
            labels = [("SUBMARKET", 0.91, 2.49, 3.0), ("2024 YEAR-END", 4.00, 2.49, 2.4),
                      ("2025 YEAR-END", 6.86, 2.49, 2.4), ("PYTHON CHANGE", 9.66, 2.49, 2.6)]
            for label, x, y, w in labels:
                text(slide, label, x, y, w, 0.32, size=10, color=MUTED, bold=True)
            for j, (area, previous, current, delta) in enumerate(data["rows"]):
                y = 2.92 + j * 1.10
                rect(slide, 0.66, y, 12.0, 0.94, WHITE, BORDER, True)
                text(slide, area, 0.90, y + 0.15, 2.9, 0.59, size=19, color=NAVY, bold=True)
                text(slide, previous, 4.0, y + 0.15, 2.4, 0.59, size=19)
                text(slide, current, 6.86, y + 0.15, 2.4, 0.59, size=19)
                text(slide, "+" + delta, 9.66, y + 0.15, 2.53, 0.59, size=19, color=TEAL, bold=True)
            rect(slide, 0.66, 5.29, 12.0, 0.75, "F7EBE2", radius=True)
            text(slide, data["warning"], 0.91, 5.40, 11.42, 0.55, size=12, color="75482D")
            text(slide, data["source_line"], 0.77, 6.35, 11.4, 0.35, size=11, color=TEAL)
            for label, url, x in (("2024 source ↗", SOURCES["2024"], 0.77),
                                  ("2025 source ↗", SOURCES["2025"], 4.3)):
                box = text(slide, label, x, 6.73, 3.25, 0.24, size=10, color=TEAL)
                box.text_frame.paragraphs[0].runs[0].hyperlink.address = url
        elif i == 4:
            for j, (num, title, body) in enumerate(data["rules"]):
                x, y = 0.66 + (j % 2) * 6.13, 2.38 + (j // 2) * 1.74
                rect(slide, x, y, 5.88, 1.52, WHITE, BORDER, True)
                text(slide, num, x + 0.22, y + 0.17, 0.64, 0.43, size=20, color=TEAL, bold=True)
                text(slide, title, x + 0.95, y + 0.15, 4.65, 0.42, size=17, color=NAVY, bold=True)
                text(slide, body, x + 0.95, y + 0.66, 4.61, 0.70, size=12)
            rect(slide, 0.66, 6.17, 12.0, 0.58, NAVY, radius=True)
            text(slide, data["bar"], 0.96, 6.29, 11.43, 0.30, size=11, color=MINT, bold=True)
        else:
            text(slide, "CURRENT SCOPE", 0.76, 2.34, 5.3, 0.37, size=12, color=TEAL, bold=True)
            text(slide, "REVIEWER ROUTE", 6.8, 2.34, 5.3, 0.37, size=12, color=TEAL, bold=True)
            for j, (label, body) in enumerate(data["scope"]):
                y = 2.82 + j * 1.02
                text(slide, label, 0.77, y, 5.4, 0.32, size=12, color=NAVY, bold=True)
                text(slide, body, 0.77, y + 0.32, 5.54, 0.72, size=12)
            for j, (num, label, body) in enumerate(data["route"]):
                y = 2.82 + j * 1.02
                text(slide, num, 6.81, y, 0.4, 0.32, size=15, color=TEAL, bold=True)
                text(slide, label, 7.23, y, 4.68, 0.32, size=12, color=NAVY, bold=True)
                text(slide, body, 7.23, y + 0.32, 5.1, 0.72, size=12)
            rect(slide, 0.66, 6.13, 12.0, 0.76, NAVY, radius=True)
            for j, (label, key) in enumerate((("Offline business tests", "offline_tests"),
                                               ("Map / E2E", "map_e2e"),
                                               ("Live model", "live_model"))):
                x = 0.95 + j * 4.05
                text(slide, label.upper(), x, 6.20, 3.8, 0.23, size=9, color=MINT, bold=True)
                text(slide, data["verification"][key], x, 6.45, 3.8, 0.27, size=12, color=WHITE, bold=True)
    prs.save(destination)


def h(value: str) -> str:
    return escape(value, quote=True)


def html_body(data: dict, i: int, screenshot_uri: str | None) -> str:
    if i == 1:
        cards = [(data["problem"], data["problem_text"]), (data["value"], data["value_text"]),
                 (data["boundary"], data["boundary_text"])]
        return '<div class="grid3">' + ''.join(f'<article class="card"><span class="index">0{j}</span><h2>{h(a)}</h2><p>{h(b)}</p></article>' for j, (a,b) in enumerate(cards, 1)) + f'</div><div class="tail">{h(data["footer"])}</div>'
    if i == 2:
        steps = ''.join(f'<div class="step"><span>{h(n)}</span><div><h2>{h(t)}</h2><p>{h(b)}</p></div></div>' for n,t,b in data["steps"])
        visual = (f'<figure><img src="{screenshot_uri}" alt="Locally supplied application screenshot"><figcaption>{h(data["image_caption"])}</figcaption></figure>'
                  if screenshot_uri else f'<div class="workflow"><strong>SELECT → ASK → CHECK</strong><p>Market context　　　　 Python tools　　　　 Source trail</p></div><small>{h(data["no_image"])}</small>')
        return f'<div class="split"><div>{steps}</div><div>{visual}</div></div>'
    if i == 3:
        rows = ''.join(f'<div class="table-row"><b>{h(a)}</b><span>{h(b)}</span><span>{h(c)}</span><strong>+{h(d)}</strong></div>' for a,b,c,d in data["rows"])
        return ('<div class="table-head"><span>SUBMARKET</span><span>2024 YEAR-END</span><span>2025 YEAR-END</span><span>PYTHON CHANGE</span></div>'
                + rows + f'<div class="caution">{h(data["warning"])}</div><div class="source">Savills reports: '
                + f'<a href="{h(SOURCES["2024"])}">2024 year-end (31 Jan 2025) ↗</a>　 '
                + f'<a href="{h(SOURCES["2025"])}">2025 year-end (30 Jan 2026) ↗</a></div>')
    if i == 4:
        cards = ''.join(f'<article class="rule"><span>{h(n)}</span><div><h2>{h(t)}</h2><p>{h(b)}</p></div></article>' for n,t,b in data["rules"])
        return f'<div class="grid2">{cards}</div><div class="darkbar">{h(data["bar"])}</div>'
    scope = ''.join(f'<div class="review-item"><b>{h(t)}</b><p>{h(b)}</p></div>' for t,b in data["scope"])
    route = ''.join(f'<div class="review-item"><b>{h(n)}　{h(t)}</b><p>{h(b)}</p></div>' for n,t,b in data["route"])
    v = ''.join(f'<div><small>{h(label)}</small><b>{h(data["verification"][key])}</b></div>' for label,key in
                (("OFFLINE BUSINESS TESTS", "offline_tests"),("MAP / E2E", "map_e2e"),("LIVE MODEL", "live_model")))
    return f'<div class="review"><div><h3>CURRENT SCOPE</h3>{scope}</div><div><h3>REVIEWER ROUTE</h3>{route}</div></div><div class="status">{v}</div>'


CSS = """
@page { size: 13.333in 7.5in; margin: 0; }
* { box-sizing: border-box; }
body { margin: 0; background: #dce4e2; font-family: Aptos, 'Helvetica Neue', Arial, sans-serif; color: #183446; }
.slide { width: 13.333in; height: 7.5in; margin: 18px auto; position: relative; overflow: hidden; background: #f7f8f5; padding: .34in .66in .45in; border-top: .12in solid #102a3a; box-shadow: 0 10px 30px #102a3a20; }
header { display: flex; justify-content: space-between; font-size: 10pt; font-weight: 700; letter-spacing: .11em; color: #206f69; }
header span:last-child { color: #617582; font-weight: 400; }
h1 { color: #102a3a; font-size: 29pt; line-height: 1.1; margin: .44in 0 .13in; letter-spacing: -.025em; }
.subtitle { color: #617582; font-size: 14pt; margin: 0; }
.content { margin-top: .35in; }
footer { position: absolute; bottom: .14in; left: .66in; right: .66in; border-top: 1px solid #dce4e2; padding-top: .06in; display: flex; justify-content: space-between; font-size: 9pt; color: #617582; letter-spacing: .07em; }
.grid3 { display: grid; grid-template-columns: repeat(3,1fr); gap: .25in; margin-top: .12in; }
.card { height: 3.37in; padding: .3in .25in; background: white; border: 1px solid #dce4e2; border-radius: 12px; }
.index { color: #206f69; font-size: 23pt; border-top: 5px solid #206f69; display: inline-block; padding-top: .18in; }
h2 { font-size: 18pt; margin: .17in 0 .1in; color: #102a3a; }
p { line-height: 1.38; }
.card p { font-size: 15pt; }
.tail { font-size: 11pt; font-weight: 700; color: #206f69; margin-top: .38in; }
.split { display:grid; grid-template-columns: 5.24in 6.1in; gap: .35in; }
.step { display:flex; gap: .24in; min-height: 1.45in; }
.step > span { color: #206f69; font-size: 22pt; font-weight: 700; }
.step h2 { font-size: 17pt; margin: 0; }
.step p { font-size: 12pt; margin: .12in 0 0; }
.workflow { background: #102a3a; border-radius: 12px; height: 3.45in; color: #b8e8d8; display: flex; align-items: center; justify-content: center; flex-direction: column; }
.workflow strong { font-size: 24pt; }
.workflow p { font-size: 13pt; color: white; }
.split small { display:block; color: #617582; margin-top: .25in; font-size: 10pt; }
figure { margin: 0; background: white; border: 1px solid #dce4e2; border-radius: 12px; height: 4.36in; padding: .15in; }
figure img { width: 100%; height: 3.55in; object-fit: contain; }
figcaption { font-size: 9pt; color: #617582; }
.table-head, .table-row { display:grid; grid-template-columns: 3.1in 2.86in 2.8in 2.7in; align-items:center; padding: 0 .24in; }
.table-head { margin-top: .16in; height: .42in; color: #617582; font-size: 10pt; font-weight: 700; }
.table-row { background: white; border: 1px solid #dce4e2; border-radius: 10px; height: .94in; margin-bottom: .16in; font-size: 19pt; }
.table-row b { color: #102a3a; }
.table-row strong { color: #206f69; }
.caution { background:#f7ebe2; color:#75482d; border-radius: 9px; padding:.14in .24in; font-size: 12pt; margin-top:.24in; }
.source { font-size: 11pt; margin-top:.3in; color:#617582; }
a { color: #206f69; text-decoration: underline; }
.grid2 { display:grid; grid-template-columns: repeat(2, 1fr); gap:.22in .25in; }
.rule { display:flex; gap:.25in; height:1.52in; padding:.17in .22in; background:white; border:1px solid #dce4e2; border-radius:10px; }
.rule > span { color:#206f69; font-size:20pt; font-weight:700; }
.rule h2 { font-size:17pt; margin:0; }
.rule p { font-size:12pt; margin:.1in 0; }
.darkbar { background:#102a3a; color:#b8e8d8; padding:.16in .28in; border-radius:9px; margin-top:.27in; font-size:11pt; font-weight:700; }
.review { display:grid; grid-template-columns:1fr 1fr; gap:.5in; }
h3 { color:#206f69; font-size:12pt; margin:.02in 0 .2in; }
.review-item { min-height: 1.02in; font-size:12pt; }
.review-item b { color:#102a3a; }
.review-item p { margin:.07in 0 0; }
.status { display:grid; grid-template-columns:repeat(3,1fr); position:absolute; left:.66in; right:.66in; bottom:.63in; background:#102a3a; color:white; border-radius:9px; padding:.09in .29in; }
.status small { display:block; color:#b8e8d8; font-size:9pt; }
.status b { display:block; font-size:12pt; margin-top:.06in; }
.notes { width: 13.333in; margin: 0 auto 25px; background: white; padding: 10px 18px; font-size: 12px; color:#617582; }
@media print { html, body { background:white; margin:0; padding:0; } .slide { margin:0; box-shadow:none; page-break-after:always; break-after:page; -webkit-print-color-adjust:exact; print-color-adjust:exact; } .slide:last-of-type { page-break-after:auto; break-after:auto; } .notes { display:none; } }
"""


def build_html(pages: list[dict], screenshot: Path | None, destination: Path):
    screenshot_uri = None
    if screenshot:
        suffix = screenshot.suffix.lower()
        mime = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".webp": "image/webp"}.get(suffix)
        if not mime:
            raise ValueError("Screenshot must be PNG, JPEG or WebP")
        screenshot_uri = f"data:{mime};base64,{base64.b64encode(screenshot.read_bytes()).decode('ascii')}"
    sections = []
    for i, item in enumerate(pages, 1):
        page = item["page"]
        sections.append(f'<section class="slide" aria-label="Slide {i} of 5">'
                        f'<header><span>LONDON / OFFICE MONITOR</span><span>{h(page.eyebrow)}</span></header>'
                        f'<h1>{h(page.title)}</h1><p class="subtitle">{h(page.subtitle)}</p>'
                        f'<div class="content">{html_body(item, i, screenshot_uri)}</div>'
                        f'<footer><span>HISTORICAL · EVIDENCE-LED · POC</span><span>{i:02d} / 05</span></footer></section>'
                        f'<details class="notes"><summary>Speaker notes and provenance · slide {i}</summary><p>{h(page.notes)}</p></details>')
    destination.write_text('<!doctype html><html lang="en"><head><meta charset="utf-8">'
                           '<meta name="viewport" content="width=device-width,initial-scale=1">'
                           '<title>London Office Monitor · Business review</title><style>' + CSS +
                           '</style></head><body>' + ''.join(sections) + '</body></html>', encoding="utf-8")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--screenshot", type=Path, help="Real local app screenshot; no UI will be fabricated if absent")
    parser.add_argument("--verification", type=Path, help="JSON statuses: offline_tests, map_e2e, live_model")
    args = parser.parse_args()
    screenshot = args.screenshot.resolve() if args.screenshot else None
    if screenshot and (not screenshot.is_file() or screenshot.suffix.lower() not in {".png", ".jpg", ".jpeg", ".webp"}):
        parser.error("--screenshot must be an existing PNG, JPEG or WebP file")
    verification = statuses(args.verification) if args.verification else statuses(None)
    pages = content(verification)
    assert len(pages) == 5
    OUT.mkdir(parents=True, exist_ok=True)
    build_ppt(pages, screenshot, OUT / "london-office-monitor.pptx")
    build_html(pages, screenshot, OUT / "london-office-monitor.html")
    print(f"Created 5-slide PPTX and HTML in {OUT}; verification: {verification}")


if __name__ == "__main__":
    main()
