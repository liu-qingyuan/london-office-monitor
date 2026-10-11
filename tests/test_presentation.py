"""Exercise the presentation CLI without overwriting reviewed delivery artifacts."""
import json
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from london_office.data import ROOT


class PresentationTests(unittest.TestCase):
    def copy_commands(self, root):
        for name in ("scripts/build_slides.py", "presentation/build_slides.py"):
            destination = root / name
            destination.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(ROOT / name, destination)

    def test_legacy_and_presentation_commands_generate_the_same_five_pages(self):
        builder = ROOT / "presentation/build_slides.py"
        self.assertTrue(builder.is_file(), "Presentation implementation belongs in presentation/")
        with tempfile.TemporaryDirectory(prefix="london-presentation-") as directory:
            root = Path(directory)
            self.copy_commands(root)
            outputs = []
            for entry in ("scripts/build_slides.py", "presentation/build_slides.py"):
                result = subprocess.run([sys.executable, str(root / entry)], cwd=ROOT,
                                        capture_output=True, text=True, timeout=10)
                self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
                self.assertIn("Created 5-slide PPTX and HTML", result.stdout)
                output = root / "artifacts/presentation"
                html = (output / "london-office-monitor.html").read_text()
                self.assertEqual(html.count('<section class="slide"'), 5)
                self.assertIn("Not yet verified", html)
                self.assertIn("£105.26", html)
                self.assertIn("£166.61", html)
                deck = Presentation(output / "london-office-monitor.pptx")
                self.assertEqual(len(deck.slides), 5)
                notes = [slide.notes_slide.notes_text_frame.text for slide in deck.slides]
                self.assertTrue(all(notes))
                texts = [[shape.text for shape in slide.shapes if shape.has_text_frame]
                         for slide in deck.slides]
                outputs.append((html, notes, texts))
            self.assertEqual(outputs[0], outputs[1])

    def test_explicit_screenshot_and_status_fixture_survive_legacy_entry(self):
        with tempfile.TemporaryDirectory(prefix="london-presentation-") as directory:
            root = Path(directory)
            self.copy_commands(root)
            # Synthetic image and display labels only, not product/model evidence.
            screenshot = root / "synthetic-image.png"
            Image.new("RGB", (8, 4), "navy").save(screenshot)
            verification = root / "synthetic-status.json"
            verification.write_text(json.dumps({"offline_tests": "Failed", "map_e2e": "Verified"}))
            result = subprocess.run([sys.executable, str(root / "scripts/build_slides.py"),
                                     "--screenshot", str(screenshot), "--verification", str(verification)],
                                    cwd=ROOT, capture_output=True, text=True, timeout=10)
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            output = root / "artifacts/presentation"
            html = (output / "london-office-monitor.html").read_text()
            self.assertIn("data:image/png;base64,", html)
            deck = Presentation(output / "london-office-monitor.pptx")
            self.assertTrue(any(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in deck.slides[1].shapes))
            labels = " ".join(shape.text for shape in deck.slides[4].shapes if shape.has_text_frame)
            self.assertIn("Failed", labels)
            self.assertIn("Verified", labels)
            self.assertIn("Not yet verified", labels)

    def test_invalid_inputs_stop_before_generating_delivery_artifacts(self):
        with tempfile.TemporaryDirectory(prefix="london-presentation-") as directory:
            root = Path(directory)
            self.copy_commands(root)
            verification = root / "invalid-status.json"
            verification.write_text(json.dumps({"live_model": "passed"}))
            cases = [(["--screenshot", str(root / "missing.png")], 2, "existing PNG"),
                     (["--verification", str(verification)], 1, "expected Not yet verified")]
            for entry in ("scripts/build_slides.py", "presentation/build_slides.py"):
                for arguments, code, message in cases:
                    with self.subTest(entry=entry, arguments=arguments):
                        result = subprocess.run([sys.executable, str(root / entry), *arguments],
                                                cwd=ROOT, capture_output=True, text=True, timeout=10)
                        self.assertEqual(result.returncode, code, result.stdout + result.stderr)
                        self.assertIn(message, result.stderr)
                        self.assertFalse((root / "artifacts").exists())


if __name__ == "__main__":
    unittest.main()
